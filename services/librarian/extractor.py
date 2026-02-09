import os
import re
import io
import boto3
from typing import Optional
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from bs4 import BeautifulSoup
from atlassian import Confluence, Jira

class ContentExtractor:
    def __init__(self):
        # Load credentials safely from environment
        self.atlassian_url = os.environ.get("ATLASSIAN_URL")
        self.atlassian_user = os.environ.get("ATLASSIAN_USER")
        self.atlassian_api_token = os.environ.get("ATLASSIAN_API_TOKEN")
        
        self._confluence = None
        self._jira = None
        self.s3_client = boto3.client('s3')

    @property
    def confluence(self):
        if not self._confluence and self.atlassian_url:
            self._confluence = Confluence(
                url=self.atlassian_url,
                username=self.atlassian_user,
                password=self.atlassian_api_token
            )
        return self._confluence

    @property
    def jira(self):
        if not self._jira and self.atlassian_url:
            self._jira = Jira(
                url=self.atlassian_url,
                username=self.atlassian_user,
                password=self.atlassian_api_token
            )
        return self._jira

    def extract(self, source: str) -> str:
        """
        Main extraction entry point. Source can be a local path, S3 path, or URL.
        """
        # 1. URL Pattern Check (Confluence / Jira)
        if source.startswith("http"):
            if "/wiki/" in source or "/pages/" in source:
                return self._extract_confluence(source)
            if "/browse/" in source:
                return self._extract_jira(source)
            return f"[URL content not supported: {source}]"

        # 2. File Extension Check
        ext = os.path.splitext(source)[1].lower()
        
        # S3 Path handling if applicable (assuming s3://bucket/key)
        if source.startswith("s3://"):
            return self._extract_from_s3(source, ext)
            
        return self._extract_file(source, ext)

    def _extract_file(self, path: str, ext: str) -> str:
        try:
            if ext == ".pdf":
                reader = PdfReader(path)
                return "\n".join([page.extract_text() for page in reader.pages])
            
            if ext == ".docx":
                doc = Document(path)
                full_text = []
                for para in doc.paragraphs:
                    full_text.append(para.text)
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            full_text.append(cell.text)
                return "\n".join(full_text)
            
            if ext == ".pptx":
                prs = Presentation(path)
                text_runs = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text_runs.append(shape.text)
                    if slide.has_notes_slide:
                        text_runs.append(slide.notes_slide.notes_text_frame.text)
                return "\n".join(text_runs)
            
            if ext == ".md" or ext == ".txt":
                with open(path, 'r', encoding='utf-8') as f:
                    return f.read()
            
            return f"[Unsupported extension: {ext}]"
        except Exception as e:
            return f"[Error extracting {path}: {str(e)}]"

    def _extract_from_s3(self, s3_path: str, ext: str) -> str:
        try:
            parts = s3_path.replace("s3://", "").split("/", 1)
            bucket = parts[0]
            key = parts[1]
            
            response = self.s3_client.get_object(Bucket=bucket, Key=key)
            content = response['Body'].read()
            file_stream = io.BytesIO(content)

            if ext == ".pdf":
                reader = PdfReader(file_stream)
                return "\n".join([page.extract_text() for page in reader.pages])
            
            if ext == ".docx":
                doc = Document(file_stream)
                return "\n".join([p.text for p in doc.paragraphs])
            
            if ext == ".pptx":
                prs = Presentation(file_stream)
                text_runs = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text_runs.append(shape.text)
                return "\n".join(text_runs)
            
            if ext == ".md" or ext == ".txt":
                return content.decode('utf-8')
                
            return f"[Unsupported S3 extension: {ext}]"
        except Exception as e:
            return f"[Error extracting from S3 {s3_path}: {str(e)}]"

    def _extract_confluence(self, url: str) -> str:
        if not self.confluence:
            return "[Confluence credentials not configured]"
        try:
            # Simple ID parsing from URL
            match = re.search(r'pageId=(\d+)', url) or re.search(r'/pages/(\d+)', url) or re.search(r'/wiki/spaces/.*?/pages/(\d+)', url)
            if not match:
                return f"[Could not parse Confluence Page ID from {url}]"
            
            page_id = match.group(1)
            page = self.confluence.get_page_by_id(page_id, expand='body.storage')
            html_content = page['body']['storage']['value']
            soup = BeautifulSoup(html_content, 'html.parser')
            return soup.get_text(separator='\n')
        except Exception as e:
            return f"[Confluence Error: {str(e)}]"

    def _extract_jira(self, url: str) -> str:
        if not self.jira:
            return "[Jira credentials not configured]"
        try:
            # Extract key like PROJ-123
            match = re.search(r'/browse/([A-Z0-9]+-\d+)', url)
            if not match:
                return f"[Could not parse Jira Issue Key from {url}]"
            
            issue_key = match.group(1)
            issue = self.jira.issue(issue_key)
            fields = issue.get('fields', {})
            
            summary = fields.get('summary', 'No Summary')
            desc = fields.get('description', 'No Description')
            status = fields.get('status', {}).get('name', 'Unknown Status')
            
            # Acceptance criteria often in a custom field, but let's check common ones or just desc
            return f"Jira Issue: {issue_key}\nSummary: {summary}\nStatus: {status}\nDescription: {desc}"
        except Exception as e:
            return f"[Jira Error: {str(e)}]"
